from pathlib import Path
import json,re,hashlib
from docx import Document
from docx.shared import Inches,Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
base=Path(__file__).resolve().parent.parent;qa=base/'qa'
source=base/'开题报告.md';s=source.read_text();bib=json.loads((base/'sources/bibliography.json').read_text())
eqs=json.loads((qa/'equations.json').read_text()); eq_index=0
page_map=json.loads((qa/'page_map.json').read_text()) if (qa/'page_map.json').exists() else {}
D=Document();sec=D.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Cm(2.3);sec.bottom_margin=Cm(2.2);sec.left_margin=Cm(2.5);sec.right_margin=Cm(2.5)
sec.header_distance=Cm(1.1);sec.footer_distance=Cm(1.1);sec.different_first_page_header_footer=True
usable=sec.page_width-sec.left_margin-sec.right_margin
styles=D.styles
for sty in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption','Header','Footer']:
 st=styles[sty];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0);st.font.italic=False
 st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Serif CJK SC')
 pp=st.paragraph_format;pp.space_after=Pt(6);pp.widow_control=True
 if sty in ['Title','Subtitle','Heading 1','Heading 2','Heading 3','Header']:
  st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK SC')
for n,size in [('Normal',12),('Title',22),('Subtitle',13),('Heading 1',16),('Heading 2',13),('Heading 3',12),('Caption',10.5),('Header',9),('Footer',9)]:styles[n].font.size=Pt(size)
normal=styles['Normal'].paragraph_format;normal.line_spacing=Pt(22);normal.first_line_indent=Pt(24);normal.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for n in ['Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption','Header','Footer']:
 f=styles[n].paragraph_format;f.first_line_indent=Pt(0);f.line_spacing=Pt(20)
 if n.startswith('Heading'):f.keep_with_next=True;f.space_before=Pt(14 if n=='Heading 1' else 10);f.space_after=Pt(7)
 for border in styles[n].element.xpath('.//w:pBdr'):border.getparent().remove(border)
styles['Caption'].paragraph_format.space_after=Pt(9)
styles['Title'].paragraph_format.line_spacing=1.15
# Language and line breaking settings.
rpr=styles['Normal'].element.get_or_add_rPr();lang=OxmlElement('w:lang');lang.set(qn('w:val'),'en-US');lang.set(qn('w:eastAsia'),'zh-CN');rpr.append(lang)
grid=sec._sectPr.find(qn('w:docGrid'))
if grid is not None:sec._sectPr.remove(grid)
for sty in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 pp=styles[sty].element.get_or_add_pPr();snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'false');pp.append(snap)
settings=D.settings.element;update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
D.core_properties.title='面向边缘智能场景的小型机械臂云边协同控制系统的设计——开题报告'
D.core_properties.subject='云端智能规划与边缘确定性执行的双系统协同方法；模拟设备 Sim2Real 路线'
D.core_properties.author='';D.core_properties.keywords='云边协同,PCSC,ETEAC,AUTO,Sim2Real,计划交接';D.core_properties.comments=''
# Header and automatic page footer.
hp=sec.header.paragraphs[0];hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT;hp.add_run('BIG-small  ·  双系统协同方法研究开题报告')
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER;fp.add_run('—  ')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');fp._p.append(f);fp.add_run('  —')
bookmark_id=1

def bookmark(p,name):
 global bookmark_id
 st=OxmlElement('w:bookmarkStart');st.set(qn('w:id'),str(bookmark_id));st.set(qn('w:name'),name)
 en=OxmlElement('w:bookmarkEnd');en.set(qn('w:id'),str(bookmark_id));bookmark_id+=1;p._p.insert(0,st);p._p.append(en)

def link(p,label,target=None,anchor=None,sup=False,size=None):
 e=OxmlElement('w:hyperlink')
 if target:e.set(qn('r:id'),p.part.relate_to(target,RT.HYPERLINK,is_external=True))
 if anchor:e.set(qn('w:anchor'),anchor)
 r=OxmlElement('w:r');rp=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'000000');rp.append(co)
 if sup:
  va=OxmlElement('w:vertAlign');va.set(qn('w:val'),'superscript');rp.append(va)
  sz=OxmlElement('w:sz');sz.set(qn('w:val'),'18');rp.append(sz)
 else:
  un=OxmlElement('w:u');un.set(qn('w:val'),'single' if target else 'none');rp.append(un)
 if size:
  sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(int(size*2)));rp.append(sz)
 r.append(rp);t=OxmlElement('w:t');t.text=label;r.append(t);e.append(r);p._p.append(e)

def write_inline(p,text):
 parts=re.split(r'(\[\d+\])',text)
 for part in parts:
  if re.fullmatch(r'\[\d+\]',part):link(p,part,anchor='ref_'+part[1:-1],sup=True)
  else:
   # Word text subscripts keep small mathematical symbols legible across CJK fonts.
   subs={'₀':'0','₁':'1','₂':'2','₃':'3','₄':'4','₅':'5','ₑ':'e','ₖ':'k','ᵢ':'i','ⱼ':'j','ₚ':'p','꜀':'c'}
   pieces=re.split(r'([₀₁₂₃₄₅ₑₖᵢⱼₚ꜀]+|T_c|τ(?:obs|gen|recv|commit))',part)
   for piece in pieces:
    if not piece:continue
    if piece=='T_c':p.add_run('T');rr=p.add_run('c');rr.font.subscript=True
    elif re.fullmatch(r'τ(?:obs|gen|recv|commit)',piece):p.add_run('τ');rr=p.add_run(piece[1:]);rr.font.subscript=True
    elif all(x in subs for x in piece):rr=p.add_run(''.join(subs[x] for x in piece));rr.font.subscript=True
    else:p.add_run(piece)

def picture(p,path,width,alt):
 run=p.add_run();img=run.add_picture(str(path),width=width);img._inline.docPr.set('descr',alt);img._inline.docPr.set('title',alt)

def table(lines,number):
 data=[[x.strip() for x in ln.strip().strip('|').split('|')] for ln in lines];data=[row for row in data if not all(re.fullmatch(r'[-: ]+',x) for x in row)]
 cap=D.add_paragraph(f'表 {number}  '+['协同策略与安全状态的正交关系','计划封装体字段与验证目的','实验指标与统计口径','研究进度与阶段成果','当前工程事实与研究任务映射'][number-1],style='Caption');cap.paragraph_format.keep_with_next=True
 tb=D.add_table(rows=1,cols=len(data[0]));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
 fractions=[.17,.22,.26,.35] if len(data[0])==4 else [.19,.38,.43]
 if number==4:fractions=[.22,.14,.30,.34]
 for c,frac in zip(tb.columns,fractions):c.width=int(usable*frac)
 for j,row in enumerate(data):
  cells=tb.rows[0].cells if j==0 else tb.add_row().cells
  trpr=tb.rows[j]._tr.get_or_add_trPr();nosplit=OxmlElement('w:cantSplit');trpr.append(nosplit)
  if j==0:rep=OxmlElement('w:tblHeader');trpr.append(rep)
  for k,txt in enumerate(row):
   cell=cells[k];cell.width=int(usable*fractions[k]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
   for edge,val in [('top',100),('bottom',100),('left',110),('right',110)]:
    node=OxmlElement('w:'+edge);node.set(qn('w:w'),str(val));node.set(qn('w:type'),'dxa');marg.append(node)
   tcpr.append(marg);borders=OxmlElement('w:tcBorders')
   for ed in ['top','left','bottom','right']:
    bd=OxmlElement('w:'+ed);bd.set(qn('w:val'),'single');bd.set(qn('w:sz'),'4');bd.set(qn('w:color'),'D9D9D9');borders.append(bd)
   tcpr.append(borders)
   if j==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEDED');tcpr.append(sh)
   p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(16)
   p.alignment=WD_ALIGN_PARAGRAPH.CENTER if (j==0 or (number==4 and k in [0,1]) or (number==5 and k==0)) else WD_ALIGN_PARAGRAPH.LEFT
   write_inline(p,txt)
   for rr in p.runs:rr.font.size=Pt(10.5);rr.bold=(j==0)
 D.add_paragraph().paragraph_format.space_after=Pt(0)

# Formal cover: only known metadata.
p=D.add_paragraph();p.paragraph_format.space_after=Pt(42)
p=D.add_paragraph('开题报告',style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(25);p.runs[0].font.size=Pt(26)
p=D.add_paragraph('面向边缘智能场景的小型机械臂\n云边协同控制系统的设计',style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(35);p.runs[0].font.size=Pt(21)
for text in ['研究重点：双系统协同的方法研究','云端智能规划系统与边缘确定性执行系统','研究路线：基于模拟设备的 Sim2Real']:
 p=D.add_paragraph(text,style='Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(10)
p=D.add_paragraph();p.paragraph_format.space_after=Pt(25)
for text in ['编制日期：2026 年 10 月 2 日','研究基础：BIG-small 当前工程及用户提供的专利文稿']:
 p=D.add_paragraph(text);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.runs[0].font.size=Pt(11)
D.add_page_break()
# Contents page is populated from final PDF page mapping.
D.add_heading('目录',level=1)
body=s[s.index('## 摘要'):].split('## 参考文献')[0]
headings=re.findall(r'^## (.+)$',body,re.M)+['参考文献']
for i,title in enumerate(headings):
 p=D.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(8);p.paragraph_format.line_spacing=Pt(22);p.paragraph_format.tab_stops.add_tab_stop(usable,WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
 link(p,title,anchor=f'sec_{i}');p.add_run('\t'+str(page_map.get(title,'待更新')))
D.add_page_break()
# Manuscript parsing.
chunks=body.split('\n\n');table_no=0;heading_no=0;captions={'architecture':'图 1  拟研究的云端慢循环与边缘快循环协同结构','handover':'图 2  新计划准备、提交前复核与有条件回退','experiment':'图 3  模拟设备配对实验与证据验证流程'}
for block in chunks:
 block=block.strip()
 if not block:continue
 if block.startswith('## '):
  title=block[3:]
  if title.startswith('一、'):D.add_page_break()
  if title.startswith('附录 A'):D.add_page_break()
  p=D.add_heading(title,level=1);bookmark(p,f'sec_{heading_no}');heading_no+=1
 elif block.startswith('### '):D.add_heading(block[4:],level=2)
 elif block.startswith(':::equation '):
  e=eqs[eq_index];eq_index+=1;p=D.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(8);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_together=True;p.paragraph_format.line_spacing=1.0
  w=min(5.85,e['width_px']/96*1.03);picture(p,qa/e['path'],Inches(w),'公式 '+str(eq_index)+': '+e['latex']);p.add_run(f'  （{eq_index}）').font.size=Pt(10)
 elif block.startswith(':::figure '):
  name=block.split()[1];p=D.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(4);p.paragraph_format.keep_with_next=True;p.paragraph_format.line_spacing=1.0;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  picture(p,qa/(name+'.png'),Inches(6.4),captions[name]);cap=D.add_paragraph(captions[name],style='Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
 elif block.startswith('|'):
  table_no+=1;table(block.splitlines(),table_no)
 else:
  p=D.add_paragraph();write_inline(p,block.replace('\n',' '))
  if block.startswith('关键词：'):p.paragraph_format.first_line_indent=Pt(0)
D.add_page_break();p=D.add_heading('参考文献',level=1);bookmark(p,f'sec_{heading_no}')
for b in bib:
 p=D.add_paragraph();p.paragraph_format.first_line_indent=Cm(-.8);p.paragraph_format.left_indent=Cm(.8);p.paragraph_format.line_spacing=Pt(16);p.paragraph_format.space_after=Pt(8);p.paragraph_format.keep_together=True;p.alignment=WD_ALIGN_PARAGRAPH.LEFT
 p.add_run(f'[{b["id"]}] '+b['text']+' ');bookmark(p,'ref_'+str(b['id']));link(p,'原始来源' if b['id']<=32 else '资料文件',target=b['url'],size=10.5)
 for r in p.runs:r.font.size=Pt(10.5)
output=base/'开题报告_双系统协同方法研究.docx';D.save(output)
report={'output':str(output),'body_cjk_count':len(re.findall('[\u4e00-\u9fff]',body)),'body_nonwhitespace_characters':len(re.sub(r'\s','',body)),'body_paragraphs':len(chunks),'references':len(bib),'external_academic_references':32,'internal_sources':2,'equations':eq_index,'figures':len(captions),'tables':table_no,'manuscript_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'source_snapshot':'396325729364d1014eb0c7beed9eb1aa11936ed7','page_map_applied':bool(page_map)}
(qa/'authoring_audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps(report,ensure_ascii=False))
