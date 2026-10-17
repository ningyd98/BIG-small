from pathlib import Path
from copy import deepcopy
from hashlib import sha256
from zipfile import ZipFile
from io import BytesIO
import json
from lxml import etree
from docx import Document
from docx.shared import Inches

ROOT = Path('/home/ningyd/文档/ChatGPT/BIGsmall')
WORK = Path('/tmp/bigsmall-system-architecture-20261011')
OUTDIR = ROOT / 'output/doc/20261024-opening-application'
SOURCE = OUTDIR / '开题申请_三项正文_20261024_计划至202704.docx'
OUTPUT = OUTDIR / '开题申请_含系统架构图_计划至202704.docx'
PNG = OUTDIR / '系统架构图_云边端.png'
CONTENT = ROOT / 'artifacts/research/process/20261004-ced-development/opening-plan-to-202704-20261011/content.json'
EXPECTED_SOURCE_SHA = '8a7168e908bd4761c42e47cba6ea61cce546713785b5bf9aac7afef30efcaf76'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
REL = 'http://schemas.openxmlformats.org/package/2006/relationships'
WP = 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
CT = 'http://schemas.openxmlformats.org/package/2006/content-types'
NS = {'w':W, 'r':R, 'wp':WP, 'a':A}
RID = 'rIdSystemArchitecture20261011'
MEDIA = 'word/media/system-architecture.png'
CAPTION = '图1 小型机械臂云边协同控制系统总体架构'

def digest(p):
    return sha256(p.read_bytes()).hexdigest()

def serialize(node):
    return etree.tostring(node, encoding='UTF-8', xml_declaration=True, standalone=True)

def element(tag, **attrs):
    e = etree.Element('{'+W+'}'+tag)
    for key,value in attrs.items():
        e.set('{'+W+'}'+key, str(value))
    return e

def text_p(pattern, value):
    p = deepcopy(pattern)
    for child in list(p):
        if child.tag != '{'+W+'}pPr':
            p.remove(child)
    run = element('r')
    old_rpr = pattern.find('w:r/w:rPr',NS)
    if old_rpr is not None:
        run.append(deepcopy(old_rpr))
    t=element('t'); t.text=value; run.append(t); p.append(run)
    return p

def center(p, picture=False):
    pr=p.find('w:pPr',NS)
    if pr is None:
        pr=element('pPr'); p.insert(0,pr)
    for tag in ['ind','jc','keepNext']:
        for n in pr.findall('w:'+tag,NS):
            pr.remove(n)
    pr.append(element('ind', firstLine=0, left=0, right=0))
    pr.append(element('jc', val='center'))
    if picture:
        pr.append(element('keepNext'))
        spacing=pr.find('w:spacing',NS)
        if spacing is None:
            spacing=element('spacing'); pr.append(spacing)
        spacing.set('{'+W+'}line','240')
        spacing.set('{'+W+'}lineRule','auto')
        spacing.set('{'+W+'}before','80')
        spacing.set('{'+W+'}after','80')

assert digest(SOURCE)==EXPECTED_SOURCE_SHA
assert PNG.exists() and PNG.stat().st_size>0
assert not OUTPUT.exists()
data=json.loads(CONTENT.read_text(encoding='utf-8'))

with ZipFile(SOURCE) as zin:
    baseline={name:zin.read(name) for name in zin.namelist()}
    root=etree.fromstring(baseline['word/document.xml'])
    body=root.find('w:body',NS)
    original_root=deepcopy(root)
    tables=body.findall('w:tbl',NS)
    assert len(tables)==3
    research_table=tables[1]
    left,right=research_table.findall('w:tr/w:tc',NS)
    paragraphs=right.findall('w:p',NS)
    assert ''.join(paragraphs[1].xpath('.//w:t/text()',namespaces=NS))==data['research'][0]
    figure_table=deepcopy(research_table)
    fleft,fright=figure_table.findall('w:tr/w:tc',NS)
    for p in fright.findall('w:p',NS):
        fright.remove(p)
    fright.append(deepcopy(paragraphs[0]))
    fright.append(deepcopy(paragraphs[1]))

    temporary=Document()
    temp_run=temporary.add_paragraph().add_run()
    shape=temp_run.add_picture(BytesIO(PNG.read_bytes()),width=Inches(5.45))
    drawing=deepcopy(shape._inline.getparent())
    blip=drawing.find('.//a:blip',NS)
    blip.set('{'+R+'}embed',RID)
    current_ids=[int(n.get('id')) for n in root.findall('.//wp:docPr',NS)]
    docpr=drawing.find('.//wp:docPr',NS)
    docpr.set('id',str(max(current_ids+[0])+1))
    docpr.set('name','系统架构图')
    docpr.set('descr','云端Qwen3.8-Max规划、边缘状态与条件判断、端侧OpenCV与RGB-D感知及受约束动作执行，通过异步请求、观测与执行反馈形成闭环。')
    picture_p=text_p(paragraphs[1],'')
    empty_run=picture_p.find('w:r',NS)
    picture_p.remove(empty_run)
    picture_run=element('r'); picture_run.append(drawing); picture_p.append(picture_run)
    center(picture_p,picture=True)
    caption_p=text_p(paragraphs[1],CAPTION)
    center(caption_p)
    fright.append(picture_p); fright.append(caption_p)

    # The source research paragraphs remain unchanged, only move to a continuation page.
    right.remove(paragraphs[0]); right.remove(paragraphs[1])
    for t in left.findall('.//w:t',NS):
        t.text=''
    old_index=list(body).index(research_table)
    prior_break=list(body)[old_index-1]
    assert prior_break.find('.//w:br',NS).get('{'+W+'}type')=='page'
    body.insert(old_index,figure_table)
    body.insert(old_index+1,deepcopy(prior_break))

    relroot=etree.fromstring(baseline['word/_rels/document.xml.rels'])
    assert RID not in [n.get('Id') for n in relroot]
    rel=etree.Element('{'+REL+'}Relationship')
    rel.set('Id',RID); rel.set('Type',R+'/image'); rel.set('Target','media/system-architecture.png'); relroot.append(rel)
    types=etree.fromstring(baseline['[Content_Types].xml'])
    types_changed=False
    if not any(n.get('Extension')=='png' for n in types):
        n=etree.Element('{'+CT+'}Default'); n.set('Extension','png'); n.set('ContentType','image/png'); types.append(n); types_changed=True

    modifications={'word/document.xml':serialize(root),'word/_rels/document.xml.rels':serialize(relroot)}
    if types_changed:
        modifications['[Content_Types].xml']=serialize(types)
    with ZipFile(OUTPUT,'w') as zout:
        for info in zin.infolist():
            zout.writestr(info,modifications.get(info.filename,baseline[info.filename]))
        zout.writestr(MEDIA,PNG.read_bytes())

with ZipFile(OUTPUT) as zout:
    changed=[n for n,v in baseline.items() if v!=zout.read(n)]
    added=[n for n in zout.namelist() if n not in baseline]
    assert set(changed)==set(modifications)
    assert added==[MEDIA]
    assert zout.read(MEDIA)==PNG.read_bytes()

new_text=root.xpath('//w:t/text()',namespaces=NS)
for group in ['background','references','research','schedule']:
    for paragraph in data[group]:
        assert new_text.count(paragraph)==1,(group,paragraph[:30])
ordered_source=[t for t in original_root.xpath('//w:t/text()',namespaces=NS) if t]
ordered_final=[t for t in new_text if t and t!=CAPTION]
assert ordered_source==ordered_final
assert new_text.count(CAPTION)==1
assert new_text.count('研究、开发方案')==1
assert etree.tostring(original_root.find('w:body/w:sectPr',NS))==etree.tostring(root.find('w:body/w:sectPr',NS))
for tag in ['tblPr','tblGrid']:
    old=[etree.tostring(n) for n in original_root.xpath('//w:'+tag,namespaces=NS)]
    new=[etree.tostring(n) for n in root.xpath('//w:'+tag,namespaces=NS)]
    assert new==[old[0],old[1],old[1],old[2]]
result={'source':str(SOURCE),'source_sha256':digest(SOURCE),'output':str(OUTPUT),'output_sha256':digest(OUTPUT),'png_sha256':digest(PNG),'changed_parts':changed,'added_parts':added,'all_original_text_once_in_order':True,'unique_architecture_caption':True,'original_section_geometry_preserved':True,'all_original_parts_outside_scope_preserved':True,'source_table_properties_and_grids_preserved':True,'expected_page_count':4}
(WORK/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(result,ensure_ascii=False,indent=2))
